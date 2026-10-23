"""Excel exporter for reconciliation results.

Produces a multi-sheet workbook with styled headers, currency formatting
and an executive summary sheet.
"""

from __future__ import annotations

import os
from datetime import datetime, timezone
from typing import Dict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# --------------------------------------------------------------------------- #
# Styling constants
# --------------------------------------------------------------------------- #

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, color="FFFFFF", size=16)
TITLE_FILL = PatternFill("solid", fgColor="1F4E78")
SECTION_RED_FILL = PatternFill("solid", fgColor="C00000")
SECTION_ORANGE_FILL = PatternFill("solid", fgColor="ED7D31")
SECTION_YELLOW_FILL = PatternFill("solid", fgColor="FFC000")
SECTION_GREEN_FILL = PatternFill("solid", fgColor="70AD47")
SECTION_FONT = Font(bold=True, color="FFFFFF", size=12)
SUBTOTAL_FONT = Font(bold=True, color="000000")
BOLD_FONT = Font(bold=True)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")
RIGHT = Alignment(horizontal="right", vertical="center")

CURRENCY_FMT = '₹#,##0'

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

CURATED_WIDTHS: Dict[str, int] = {
    "AWB_Number": 18, "Marketplace": 12, "SKU": 12, "Product_Name": 28,
    "Qty": 6, "Order_Date": 12, "Order_Status": 24,
    "Return_Initiated_Date": 20, "Delivered_To_Seller_Date": 22,
    "Warehouse_Physical_Received": 20, "Zone": 10, "Courier": 14,
    "Dead_Weight_g": 14, "Length_cm": 10, "Breadth_cm": 12, "Height_cm": 10,
    "Cost_Price": 12, "Selling_Price": 12,
    "Invoice_Value": 14, "Marketplace_Fee": 14,
    "Forward_Shipping_Charged": 18, "Reverse_Shipping_Charged": 18,
    "Actual_Bank_Settlement": 18, "Claim_Settlement_Amount": 20,
    "Claim_Paid": 10, "Weight_Charged_by_Courier_g": 22,
    "Volume_Weight_g": 14, "Billable_Weight_g": 16,
    "Expected_Shipping_Fee": 16, "Claimable_Amount": 16,
    "Days_In_Transit": 14, "Recon_Type": 18,
}

CURRENCY_COLUMNS = {
    "Cost_Price", "Selling_Price", "Invoice_Value", "Marketplace_Fee",
    "Forward_Shipping_Charged", "Reverse_Shipping_Charged",
    "Actual_Bank_Settlement", "Claim_Settlement_Amount",
    "Expected_Shipping_Fee", "Claimable_Amount",
}


# --------------------------------------------------------------------------- #
# Helpers
# --------------------------------------------------------------------------- #


def _write_df_sheet(wb: Workbook, title: str, df) -> None:
    ws = wb.create_sheet(title=title)
    if df is None or len(df) == 0:
        ws["A1"] = "No rows."
        ws["A1"].font = BOLD_FONT
        ws.freeze_panes = "A2"
        return

    columns = list(df.columns)
    ws.append(columns)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER

    for row in df.itertuples(index=False):
        ws.append(list(row))

    ws.freeze_panes = "A2"

    for col_idx, col_name in enumerate(columns, start=1):
        letter = get_column_letter(col_idx)
        width = CURATED_WIDTHS.get(col_name, 15)
        ws.column_dimensions[letter].width = width
        if col_name in CURRENCY_COLUMNS:
            for r in range(2, ws.max_row + 1):
                ws.cell(row=r, column=col_idx).number_format = CURRENCY_FMT
                ws.cell(row=r, column=col_idx).alignment = RIGHT
        else:
            for r in range(2, ws.max_row + 1):
                cell = ws.cell(row=r, column=col_idx)
                cell.alignment = LEFT


def _shade_range(ws, cell_range: str, fill: PatternFill, font: Font) -> None:
    for row in ws[cell_range]:
        for cell in row:
            cell.fill = fill
            cell.font = font


# --------------------------------------------------------------------------- #
# Executive summary
# --------------------------------------------------------------------------- #


def _write_summary_sheet(wb: Workbook, summary: Dict[str, object]) -> None:
    ws = wb.create_sheet(title="Executive_Summary")

    ws.merge_cells("A1:D1")
    ws["A1"] = "E-Commerce Claim Recovery — Executive Summary"
    ws["A1"].font = TITLE_FONT
    ws["A1"].fill = TITLE_FILL
    ws["A1"].alignment = CENTER
    ws.row_dimensions[1].height = 32

    ws.merge_cells("A2:D2")
    ws["A2"] = f"Generated at: {summary.get('generated_at', datetime.now(timezone.utc).isoformat())}"
    ws["A2"].alignment = CENTER
    ws["A2"].font = Font(italic=True, color="595959")

    row = 4
    def _row_label(ws, r, label, value, value_fmt=None):
        ws.cell(row=r, column=1, value=label).font = BOLD_FONT
        c = ws.cell(row=r, column=2, value=value)
        c.alignment = RIGHT
        if value_fmt:
            c.number_format = value_fmt

    _row_label(ws, row, "Total Orders", int(summary["total_orders"]))
    row += 1
    _row_label(ws, row, "Total Recoverable", int(summary["total_recoverable"]), CURRENCY_FMT)
    row += 2

    def _section(r, title, count, amount, fill):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        ws.cell(row=r, column=1, value=title)
        for c in ws[r]:
            c.fill = fill
            c.font = SECTION_FONT
            c.alignment = LEFT
        ws.cell(row=r + 1, column=1, value="Count").font = BOLD_FONT
        ws.cell(row=r + 1, column=2, value=int(count)).alignment = RIGHT
        ws.cell(row=r + 2, column=1, value="Total Claimable").font = BOLD_FONT
        money = ws.cell(row=r + 2, column=2, value=int(amount))
        money.number_format = CURRENCY_FMT
        money.alignment = RIGHT

    _section(row, "Weight Fraud (Courier over-charge)",
             summary["weight_fraud_count"], summary["weight_fraud_amount"], SECTION_RED_FILL)
    row += 4
    _section(row, "Lost RTO (45+ days unreconciled)",
             summary["lost_rto_count"], summary["lost_rto_amount"], SECTION_ORANGE_FILL)
    row += 4
    _section(row, "Unpaid Damaged Returns",
             summary["unpaid_damaged_count"], summary["unpaid_damaged_amount"], SECTION_YELLOW_FILL)
    row += 4
    _section(row, "Clean (no claim)",
             summary["clean_count"], 0, SECTION_GREEN_FILL)

    # Column widths
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 6
    ws.column_dimensions["D"].width = 6


# --------------------------------------------------------------------------- #
# Public API
# --------------------------------------------------------------------------- #


def export_report(recon_result: dict, output_path: str = "output/claim_recovery_report.xlsx") -> str:
    """Export a multi-sheet Excel workbook from a reconciliation result.

    Args:
        recon_result: Dict returned by :func:`core.recon_engine.run_reconciliation`.
            Expected keys: ``summary``, ``weight_fraud``, ``lost_rto``,
            ``unpaid_damaged``.
        output_path: Destination xlsx path. Parent directories are created
            automatically.

    Returns:
        The resolved absolute path to the workbook.
    """
    out_dir = os.path.dirname(os.path.abspath(output_path))
    os.makedirs(out_dir, exist_ok=True)

    wb = Workbook()
    # Remove default sheet; we'll add our own.
    default = wb.active
    wb.remove(default)

    _write_summary_sheet(wb, recon_result["summary"])
    _write_df_sheet(wb, "Weight_Fraud", recon_result["weight_fraud"])
    _write_df_sheet(wb, "Lost_RTO", recon_result["lost_rto"])
    _write_df_sheet(wb, "Unpaid_Damaged", recon_result["unpaid_damaged"])

    wb.save(output_path)
    return os.path.abspath(output_path)
